# -*- coding: utf-8 -*-
"""通用文件工具类：文档文本解析 + 结果 DTO，供多模块共享（工作流 DOC_EXTRACTOR、知识库文档等）。

存储能力已抽象到 common.common_storage（StorageBackend：本地 / MinIO 等 S3 对象存储），
本类只做与存储无关的：文件内容解析（extract_text / load_and_extract）与 DTO 组装。
调用方先经存储后端取得本地路径（get_local_path），再交给本类解析。
"""
from __future__ import annotations

import asyncio
import os
from datetime import datetime
from typing import Optional


class FileUtils:
    """文档文本解析与结果 DTO 静态工具类（无业务 DB 依赖，多模块可直接复用）。"""

    # ==================== DTO ====================

    @staticmethod
    def dto(file_id: str, name: str, size: int, content_type: str,
            upload_time: Optional[datetime] = None) -> dict:
        """上传结果/文件信息 DTO（对齐前端文件选择器结构）。"""
        return {"fileId": file_id, "name": name, "size": size,
                "contentType": content_type,
                "uploadTime": (upload_time or datetime.now()).strftime("%Y-%m-%d %H:%M:%S")}

    # ==================== 文本解析 ====================

    @staticmethod
    async def load_and_extract(path: str, file_id: Optional[str] = None,
                               supported_types: Optional[list] = None) -> tuple[str, dict]:
        """本地路径 → (文本内容, 元数据)。file_id 用于从存储名还原原始文件名。

        文件引用解析（fileId → 本地路径）由各模块自行负责：工作流走
        WorkflowFileService.get_local_path，知识库可走自己的文件表。
        支持格式：txt/md/log/csv/json → 内置读取；pdf → pypdf；docx → python-docx；
        xlsx → openpyxl；xls → xlrd；pptx → python-pptx；html → 内置去标签；
        老格式 doc/ppt → LibreOffice 转换。
        """
        if not path or not os.path.exists(path):
            raise ValueError(f"文件不存在或不可访问: {path}")
        ext = os.path.splitext(path)[1].lower().lstrip(".")
        if supported_types and ext not in {str(t).lower().lstrip(".") for t in supported_types}:
            raise ValueError(f"不支持的文件类型 .{ext}（允许: {', '.join(map(str, supported_types))}）")
        try:
            content = await asyncio.to_thread(FileUtils.extract_text, path, ext)
        except ValueError:
            raise
        except Exception as e:  # noqa: BLE001
            raise ValueError(f"文档解析失败: {e}") from e
        # 元数据：存储名 = {file_id}_{原文件名}，精确剥离前缀还原原名
        # （startswith 兼容 file_id 为 None 的本地路径直用场景，不会残留连接符下划线）
        base_name = os.path.basename(path)
        stored_prefix = f"{file_id}_"
        file_name = base_name[len(stored_prefix):] if file_id and base_name.startswith(stored_prefix) else base_name
        return content, {"fileid": file_id or path, "name": file_name,
                         "size": os.path.getsize(path), "type": ext or "txt"}

    @staticmethod
    def extract_text(path: str, ext: str) -> str:
        """按扩展名提取文本（同步方法，pdf 无内嵌文本页返回空串，不报错）。"""
        if ext in ("txt", "md", "log", "csv", "json"):
            for enc in ("utf-8", "gbk"):
                try:
                    with open(path, "r", encoding=enc) as f:
                        return f.read(2 * 1024 * 1024)
                except UnicodeDecodeError:
                    continue
            with open(path, "rb") as f:
                return f.read().decode("utf-8", errors="ignore")
        if ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if ext == "docx":
            from docx import Document
            doc = Document(path)
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        if ext == "xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        cells = ["" if v is None else str(v) for v in row]
                        if any(c.strip() for c in cells):
                            rows.append("\t".join(cells))
                return "\n".join(rows)
            finally:
                wb.close()
        if ext == "xls":
            import xlrd
            book = xlrd.open_workbook(path)
            rows = []
            for sheet in book.sheets():
                for r in range(sheet.nrows):
                    cells = ["" if v is None else str(v) for v in sheet.row_values(r)]
                    if any(c.strip() for c in cells):
                        rows.append("\t".join(cells))
            return "\n".join(rows)
        if ext == "pptx":
            from pptx import Presentation
            prs = Presentation(path)
            parts = []
            for slide in prs.slides:
                texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            text = "".join(run.text for run in para.runs).strip()
                            if text:
                                texts.append(text)
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            cells = [c.text.strip() for c in row.cells]
                            if any(cells):
                                texts.append(" | ".join(cells))
                if texts:
                    parts.append("\n\n".join(texts))
            return "\n\n".join(parts)
        if ext == "html":
            import html as html_mod
            import re
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                raw = f.read(2 * 1024 * 1024)
            raw = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", raw)
            text = re.sub(r"(?s)<[^>]+>", " ", raw)
            return html_mod.unescape(re.sub(r"[ \t]+", " ", text)).strip()
        if ext in ("doc", "ppt"):
            return FileUtils._convert_via_libreoffice(path, ext)
        raise ValueError(f"不支持的文件类型: .{ext}")

    @staticmethod
    def _convert_via_libreoffice(path: str, ext: str) -> str:
        """老格式 .doc/.ppt 无纯 Python 解析库，走 LibreOffice headless 转文本。"""
        import shutil
        import subprocess
        import tempfile
        candidates = [
            shutil.which("soffice") or shutil.which("libreoffice"),
            "C:/Program Files/LibreOffice/program/soffice.exe",
            "C:/Program Files (x86)/LibreOffice/program/soffice.exe",
            "/usr/bin/soffice",
            "/usr/lib/libreoffice/program/soffice",
        ]
        soffice = next((c for c in candidates if c and os.path.exists(c)), None)
        if not soffice:
            raise ValueError(
                f"解析 .{ext} 需要安装 LibreOffice（未检测到 soffice）："
                "请安装 LibreOffice，或将文件另存为 .docx/.pptx 后重试")
        with tempfile.TemporaryDirectory(prefix="wf_extract_") as out_dir:
            try:
                proc = subprocess.run(
                    [soffice, "--headless", "--convert-to", "txt:Text (encoded):UTF8",
                     "--outdir", out_dir, path],
                    capture_output=True, timeout=120)
            except subprocess.TimeoutExpired:
                raise ValueError(f"LibreOffice 转换超时: {path}") from None
            if proc.returncode != 0:
                detail = proc.stderr.decode("utf-8", errors="ignore")[:200]
                raise ValueError(f"LibreOffice 转换失败: {detail}")
            txts = [f for f in os.listdir(out_dir) if f.lower().endswith(".txt")]
            if not txts:
                raise ValueError(f"LibreOffice 转换未产出文本: {path}")
            txt_path = os.path.join(out_dir, txts[0])
            for enc in ("utf-8", "gbk"):
                try:
                    with open(txt_path, "r", encoding=enc) as f:
                        return f.read(2 * 1024 * 1024)
                except UnicodeDecodeError:
                    continue
            with open(txt_path, "rb") as f:
                return f.read().decode("utf-8", errors="ignore")