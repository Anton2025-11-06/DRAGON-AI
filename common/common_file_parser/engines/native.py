# -*- coding: utf-8 -*-
"""native 原生解析引擎：pypdf / python-docx / python-pptx / openpyxl / xlrd 轻量直读。

特点（SPEC §7.2 第 1 项）：无模型依赖、速度快、CPU 运行，是默认引擎。
明确的短板：不做版面分析（双栏 PDF 会按流顺序串行读出）、不做复杂表格还原、扫描件抽不出字。
这三类文档该用 docling 或 minerU——auto 引擎正是按这些短板做判据的；native 只在 warnings
里如实说明，不假装成功（用户看到「解析完成但一个字没有」比看到失败更难排查）。
"""
from __future__ import annotations

import csv
import io
import json
import re
from typing import Any, Optional, Sequence

from common.common_constants import rag_constant as RC
from common.common_file_parser import constants as C
from common.common_file_parser.engines.base import (BaseEngine, ParseError,
                                                    guard_legacy_office,
                                                    normalize_ext)
from common.common_file_parser.engines.md_block_builder import (html_to_blocks,
                                                                markdown_to_blocks)
from common.common_file_parser.models import ImageRef, ParsedBlock, ParsedDocument
from common.common_file_parser.utils.text_utils import decode_bytes, normalize_space

# 单个表格块的目标字数：超过它按行拆块，避免「一个块远超单块上限」被硬截成两半（拆块保表头）
TABLE_BLOCK_CHARS = 6000
# 单表最多读取行数：read_only 模式也是逐行解析，几百万行的表会把 worker 内存吃穿
MAX_ROWS_PER_SHEET = 20000
# docx 标题样式名（中英文都认，Word 的样式名跟界面语言绑定）
_DOCX_HEADING = re.compile(r"(?:heading|标题|题注)\s*([1-6])", re.I)
# 老式 VML 命名空间：Word 2003 兼容对象/剪贴画用 <v:imagedata r:id> 引图，不在 DrawingML 里
_VML_NS = "urn:schemas-microsoft-com:vml"


def _note_image_cap(doc: ParsedDocument) -> None:
    """图片取满上限的说明：只记一句（几百页文档不该刷几百条）。"""
    msg = f"内嵌图片已取满 {C.MAX_EMBED_IMAGES} 张上限，其余图片未提取（内存保护）"
    if not any(msg in w for w in doc.warnings):
        doc.warnings.append(msg)


def _cell(value: Any) -> str:
    """单元格 → 单行安全文本（压缩空白、转义竖线，防止撑破 markdown 表格）。"""
    text = "" if value is None else str(value)
    return re.sub(r"\s+", " ", text).replace("|", "\\|").strip()


def _rows_to_markdown(rows: Sequence[Sequence[Any]], *,
                      header: Optional[Sequence[Any]] = None) -> str:
    """二维表 → markdown 表格文本（全空行丢掉；给了 header 则全部行都算数据行）。"""
    body = [[_cell(c) for c in r] for r in rows if r and any(_cell(c) for c in r)]
    if not body and not header:
        return ""
    if header is not None:
        head = [_cell(c) for c in header]
        data = body
    else:
        if not body:
            return ""
        head, data = body[0], body[1:]
    width = max([len(head)] + [len(r) for r in data])
    head = head + [""] * (width - len(head))
    lines = ["| " + " | ".join(head) + " |", "|" + "---|" * width]
    for r in data:
        r = list(r) + [""] * (width - len(r))
        lines.append("| " + " | ".join(r) + " |")
    return "\n".join(lines)


class NativeEngine(BaseEngine):
    """默认引擎：按扩展名分派到各轻量库的同步解析函数，全部在线程池里执行。"""

    name = C.ENGINE_NATIVE
    # 只声明「文档格式」：doc 型库也收图片/音频/视频，但那些不进解析引擎（由大模型先转文字，
    # 见 parse_service 的媒体转写分支）。写成 KB_TYPE_ALLOWED_EXTS[doc] 等于声称自己能读 mp4。
    # 剔掉老版 Office：base.guard_legacy_office 对本引擎一律拒收，声明里留着 doc/ppt
    # 就是让配置页照着一份假清单告诉用户「native 能读 .doc」（与 docling 同一口径）。
    formats = set(RC.DOC_TEXT_EXTS) - C.LEGACY_OFFICE

    async def _parse(self, raw: bytes, filename: str, ext: str) -> ParsedDocument:
        ext = ext or normalize_ext(filename)
        # 文案收到 base：三处引擎各自写一句，改口径时必有人漏（docling 就是这么漏的）
        guard_legacy_office(ext)
        if not raw:
            raise ParseError("文件内容为空")
        handler = {
            "pdf": self._pdf, "docx": self._docx, "pptx": self._pptx,
            "xlsx": self._xlsx, "xls": self._xls, "csv": self._csv,
            "md": self._markdown, "markdown": self._markdown, "html": self._html,
            "htm": self._html, "json": self._json, "txt": self._txt,
        }.get(ext)
        if handler is None:
            raise ParseError(f"不支持的文件格式: .{ext}（doc 型知识库支持 {sorted(self.formats)}）")
        doc = self.new_document(filename, ext)
        doc.blocks, doc.meta = await self.run_cpu(handler, raw, doc)
        return doc

    # ==================== PDF ====================

    def _pdf(self, raw: bytes, doc: ParsedDocument) -> tuple[list[ParsedBlock], dict]:
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(raw))
        except Exception as e:                          # noqa: BLE001 - 损坏文件的异常种类很杂
            raise ParseError(f"PDF 无法打开（可能已损坏或加密）: {e}") from e
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):              # 只设了权限密码的空口令 PDF
                    raise ValueError
            except (TypeError, ValueError):
                raise ParseError("PDF 已加密，无法提取文本") from None

        blocks: list[ParsedBlock] = []
        image_total = 0
        text_total = 0
        for pno, page in enumerate(reader.pages, start=1):
            try:
                page_text = page.extract_text() or ""
            except Exception as e:                      # noqa: BLE001
                doc.warnings.append(f"第 {pno} 页文本提取失败: {e}")
                page_text = ""
            text_total += len(re.sub(r"\s+", "", page_text))
            if page_text.strip():
                blocks.extend(markdown_to_blocks(page_text, page=pno))
            # 内嵌图片一律提取（要传存储并把地址回填到正文原位），只按内存上限截断
            if image_total < C.MAX_EMBED_IMAGES:
                for img in self._pdf_images(page, pno, doc):
                    image_total += 1
                    blocks.append(img)
                    if image_total >= C.MAX_EMBED_IMAGES:
                        _note_image_cap(doc)
                        break
        pages = len(reader.pages)
        if image_total >= C.MAX_EMBED_IMAGES:
            _note_image_cap(doc)
        scanned = pages > 0 and text_total / pages < C.SCANNED_MIN_CHARS_PER_PAGE
        meta = {"pages": pages, "chars": text_total, "images": image_total, "scanned": scanned}
        if scanned:
            doc.warnings.append(
                "每页可提取文字极少，判断为扫描件/图片型 PDF："
                "native 识别不了其中的图形文字，建议改用 docling 或 minerU 重新解析")
        if not blocks:
            raise ParseError("PDF 没有可提取的文本（扫描件或纯图片文档）")
        return blocks, meta

    def _pdf_images(self, page, pno: int, doc: ParsedDocument) -> list[ParsedBlock]:
        """抽内嵌图片（pypdf 需要 Pillow 解码；EMF/WMF 矢量图 pypdf 解不动，直接跳过）。

        逐图隔离：pypdf 的 page.images 是惰性的，**取下标就会解码那一张**，一张坏图抛的异常
        会穿破整次迭代——拿 list() 一次解完等于「一张图解不开就丢一整页的图」。所以先只数个数，
        再一张一张取，坏图单独记一句警告跳过（解不开是常态：JPX2000/CCITT/内联图都可能失手）。
        """
        out: list[ParsedBlock] = []
        try:
            images = page.images
            total = len(images)                         # 只取个数，不触发解码
        except Exception as e:                          # noqa: BLE001
            # 缺依赖/解不动是整篇的性质，不是某一页的偶发事：只记一句，免得几百页刷几百条
            if not any("图片提取失败" in w for w in doc.warnings):
                doc.warnings.append(f"内嵌图片提取失败（第 {pno} 页起）: {e}")
            return out
        dropped = 0
        for idx in range(total):
            try:
                img = images[idx]
                name = str(getattr(img, "name", "") or getattr(img, "imagename", "") or "")
                data = getattr(img, "data", None)
            except Exception as e:                      # noqa: BLE001 - 单张坏图不许带走整页
                dropped += 1
                if dropped <= 3:
                    doc.warnings.append(f"第 {pno} 页第 {idx + 1} 张内嵌图片解不开，已跳过: {e}")
                continue
            ext = normalize_ext(name)
            if ext in ("emf", "wmf"):
                continue
            if not data:
                continue
            ref = ImageRef(page=pno, mime=f"image/{ext or 'png'}", data=bytes(data))
            out.append(ParsedBlock(type=C.BLOCK_IMAGE, text="", page=pno, image=ref))
        if dropped > 3:
            doc.warnings.append(f"第 {pno} 页共 {dropped} 张内嵌图片解不开（以上只列前 3 张）")
        return out

    # ==================== DOCX ====================

    def _docx(self, raw: bytes, doc: ParsedDocument) -> tuple[list[ParsedBlock], dict]:
        import docx
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        try:
            d = docx.Document(io.BytesIO(raw))
        except Exception as e:                          # noqa: BLE001
            raise ParseError(f"DOCX 无法打开（不是标准的 Office Open XML）: {e}") from e

        blocks: list[ParsedBlock] = []
        chars = 0
        images = 0
        seen_rids: set[str] = set()      # 同一张图被引多次只抽一份（页眉 logo / AlternateContent 副本）
        # 按 body 子元素顺序遍历：段落与表格交错时保留原始顺序（SPEC §7.4 结构化还原）
        for child in d.element.body.iterchildren():
            if child.tag == qn("w:p"):
                para = Paragraph(child, d)
                text = normalize_space(para.text)
                level = 0
                style_name = str(getattr(para.style, "name", "") or "")
                hit = _DOCX_HEADING.search(style_name)
                if hit:
                    level = int(hit.group(1))
                elif style_name.lower() in ("title", "文档标题"):
                    level = 1
                if text:
                    chars += len(text)
                    blocks.append(ParsedBlock(type=C.BLOCK_TITLE if level else C.BLOCK_TEXT,
                                              text=text, level=level))
                # 图文顺序：先落文字再落图，与原文版面一致（图排在前面会让“如下图所示”找不到指代）
                for ref in self._docx_images(d, child, qn, seen_rids):
                    if images >= C.MAX_EMBED_IMAGES:
                        _note_image_cap(doc)
                        break
                    images += 1
                    blocks.append(ParsedBlock(type=C.BLOCK_IMAGE, text="", image=ref))
            elif child.tag == qn("w:tbl"):
                md = _rows_to_markdown([[c.text for c in row.cells]
                                        for row in Table(child, d).rows])
                if md:
                    chars += len(md)
                    blocks.append(ParsedBlock(type=C.BLOCK_TABLE, text=md))
                # 表格单元格（含单元格里嵌的文本框）内嵌的图同样靠 a:blip 引图，
                # 只走上面段落分支会一整张表里一张图都抽不出
                for ref in self._docx_images(d, child, qn, seen_rids):
                    if images >= C.MAX_EMBED_IMAGES:
                        _note_image_cap(doc)
                        break
                    images += 1
                    blocks.append(ParsedBlock(type=C.BLOCK_IMAGE, text="", image=ref))
        if not blocks:
            raise ParseError("DOCX 里没有可提取的正文")
        return blocks, {"pages": 0, "chars": chars, "images": images}

    def _docx_images(self, d, element, qn, seen: Optional[set] = None) -> list[ImageRef]:
        """元素子树里嵌入的媒体（图片/音频/视频，含表格单元格与文本框里的）：按关系 ID 取部件字节。

        不按 mime 过滤：Word 把内嵌音频/视频也当媒体部件存在 word/media/ 下、由同一个 a:blip
        引用，content_type 是 video/mp4、audio/mpeg——只认 image/ 开头就会把它们整批丢掉。
        种类交给 ImageRef.kind 按 mime 认，上层据此决定调哪个模型做解析增强。

        三个容易漏的点：
          - Word 2003 兼容对象、部分剪贴画/艺术字走的是老式 VML（<v:imagedata r:id>），
            只认 a:blip 会整张图看不见；
          - mc:AlternateContent 的 Choice 与 Fallback 各存一份同样的引用，不去重就会上传两遍；
          - 走 <o:OLEObject> 嵌入的媒体拿到的是 OLE 包（application/vnd…oleObject），
            不是媒体字节，按图片传上去只会多一个打不开的对象——这里不认它。
        """
        out: list[ImageRef] = []
        rels = d.part.related_parts
        seen = seen if seen is not None else set()
        for node in element.iter():
            tag = str(getattr(node, "tag", "") or "")
            rid = ""
            if tag.endswith("}blip") and "drawingml" in tag:
                rid = node.get(qn("r:embed")) or node.get(qn("r:link")) or ""
            elif tag.endswith("}imagedata") and _VML_NS in tag:
                rid = node.get(qn("r:id")) or ""
            if not rid or rid in seen:
                continue
            part = rels.get(rid)
            data = getattr(part, "blob", None)
            if not data:
                continue
            seen.add(rid)
            out.append(ImageRef(mime=str(getattr(part, "content_type", "") or "image/png"),
                                data=bytes(data)))
        return out

    # ==================== PPTX ====================

    def _pptx(self, raw: bytes, doc: ParsedDocument) -> tuple[list[ParsedBlock], dict]:
        from pptx import Presentation

        try:
            prs = Presentation(io.BytesIO(raw))
        except Exception as e:                          # noqa: BLE001
            raise ParseError(f"PPTX 无法打开: {e}") from e

        blocks: list[ParsedBlock] = []
        chars = 0
        images = 0
        total = 0
        for slide in prs.slides:
            total += 1
            slide_blocks: list[ParsedBlock] = []
            title_text = ""
            for shape in slide.shapes:
                if getattr(shape, "has_table", False):
                    md = _rows_to_markdown([[c.text for c in row.cells]
                                            for row in shape.table.rows])
                    if md:
                        chars += len(md)
                        slide_blocks.append(ParsedBlock(type=C.BLOCK_TABLE, text=md, page=total))
                    continue
                if getattr(shape, "has_text_frame", False):
                    text = normalize_space(
                        "\n".join(p.text for p in shape.text_frame.paragraphs))
                    if not text:
                        continue
                    if not title_text and getattr(shape, "is_title", False):
                        title_text = text      # 标题单独成块，供按页/按标题分块定位
                        continue
                    chars += len(text)
                    slide_blocks.append(ParsedBlock(type=C.BLOCK_TEXT, text=text, page=total,
                                                     meta={"slide": total}))
                    continue
                ref = self._pptx_image(shape)
                if ref is not None:
                    if images >= C.MAX_EMBED_IMAGES:
                        _note_image_cap(doc)
                        continue
                    images += 1
                    slide_blocks.append(ParsedBlock(type=C.BLOCK_IMAGE, text="", page=total,
                                                     image=ref))
            if title_text:
                blocks.append(ParsedBlock(type=C.BLOCK_TITLE, text=title_text, level=2,
                                          page=total))
            blocks.extend(slide_blocks)
        if not blocks:
            raise ParseError("PPT 里没有可提取的文本")
        return blocks, {"pages": total, "chars": chars, "images": images}

    def _pptx_image(self, shape) -> Optional[ImageRef]:
        """形状里的图片/音视频（一律提取）；两种都不像的 shape 返回 None。

        顺序不能倒：PPT 里插入的视频/音频在 python-pptx 上是 media shape，真字节在
        ``shape.media_file`` 里，而 ``shape.image`` 读到的是那张封面图——先试 image 就会
        把「一段视频」解析成「一张海报」。非媒体 shape 访问 media_file 会抛 TypeError，
        旧版 python-pptx 根本没有这个属性（AttributeError），两者都退到图片分支。
        """
        blob: bytes = b""
        mime = ""
        try:
            media_file = shape.media_file
            blob = bytes(media_file.blob)
            mime = str(getattr(media_file, "content_type", "") or "")
        except Exception:                               # noqa: BLE001 - 非媒体形状，退图片
            try:
                blob = bytes(shape.image.blob)
                mime = f"image/{getattr(shape.image, 'ext', '') or 'png'}"
            except Exception:                           # noqa: BLE001 - 非图片形状没有 image
                return None
        if not blob:
            return None
        return ImageRef(mime=mime or "image/png",
                        alt=str(getattr(shape, "name", "") or ""), data=bytes(blob))

    # ==================== Excel / CSV ====================

    def _table_blocks(self, sheet: str, rows: list[list[str]],
                      doc: ParsedDocument) -> list[ParsedBlock]:
        """数据行 → 若干表格块（首行作表头，每块都自带表头，字数超上限就换块）。

        这里就把表头重复写进每一块：excel 策略（SPEC §7.5）要求「每块强制附带原始表头」，
        而 sidecar 只存块内容——分块阶段再想补表头就没有表头原文了。
        """
        if not rows:
            return []
        header, body = rows[0], rows[1:]
        blocks: list[ParsedBlock] = []
        chunk: list[list[str]] = []
        flush_at = len(body)
        for i, row in enumerate(body):
            chunk.append(row)
            if len(_rows_to_markdown(chunk, header=header)) >= TABLE_BLOCK_CHARS or \
                    i == flush_at - 1:
                md = _rows_to_markdown(chunk, header=header)
                if md:
                    blocks.append(self._sheet_block(sheet, md))
                chunk = []
        if chunk:
            blocks.append(self._sheet_block(sheet, _rows_to_markdown(chunk, header=header)))
        if not blocks:                                  # 只有表头没有数据行
            blocks.append(self._sheet_block(sheet, _rows_to_markdown([], header=header)))
            doc.warnings.append(f"工作表 {sheet} 只有表头、没有数据行")
        return blocks

    def _sheet_block(self, sheet: str, md: str) -> ParsedBlock:
        lines = (md or "").split("\n")
        return ParsedBlock(type=C.BLOCK_TABLE, text=md, sheet=sheet,
                           meta={"header": lines[0] if lines else "",
                                 "rows": max(0, len(lines) - 2)})

    def _xlsx(self, raw: bytes, doc: ParsedDocument) -> tuple[list[ParsedBlock], dict]:
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as e:                          # noqa: BLE001
            raise ParseError(f"XLSX 无法打开: {e}") from e
        blocks: list[ParsedBlock] = []
        chars = 0
        sheet_total = 0
        try:
            sheet_total = len(wb.sheetnames)
            for ws in wb.worksheets:
                rows: list[list[str]] = []
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= MAX_ROWS_PER_SHEET:
                        doc.warnings.append(
                            f"工作表 {ws.title} 超过 {MAX_ROWS_PER_SHEET} 行，只解析前面部分")
                        break
                    cells = [_cell(v) for v in row]
                    if any(cells):
                        rows.append(cells)
                piece = self._table_blocks(ws.title, rows, doc)
                chars += sum(len(b.text) for b in piece)
                blocks.extend(piece)
        finally:
            wb.close()
        if not blocks:
            raise ParseError("Excel 里没有数据行")
        return blocks, {"pages": 0, "chars": chars, "images": 0, "sheets": sheet_total}

    def _xls(self, raw: bytes, doc: ParsedDocument) -> tuple[list[ParsedBlock], dict]:
        try:
            import xlrd
        except ImportError as e:                        # pragma: no cover
            raise ParseError("读取 .xls 需要 xlrd 依赖（pip install xlrd），"
                             "或先把文件另存为 .xlsx") from e
        try:
            book = xlrd.open_workbook(file_contents=raw)
        except Exception as e:                          # noqa: BLE001
            raise ParseError(f"XLS 无法打开: {e}") from e
        blocks: list[ParsedBlock] = []
        chars = 0
        for sheet in book.sheets():
            rows = [[_cell(v) for v in sheet.row_values(r)] for r in range(sheet.nrows)]
            rows = [r for r in rows if any(r)]
            piece = self._table_blocks(sheet.name, rows, doc)
            chars += sum(len(b.text) for b in piece)
            blocks.extend(piece)
        if not blocks:
            raise ParseError("XLS 里没有数据行")
        return blocks, {"pages": 0, "chars": chars, "images": 0, "sheets": book.nsheets}

    def _csv(self, raw: bytes, doc: ParsedDocument) -> tuple[list[ParsedBlock], dict]:
        text = decode_bytes(raw)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel                        # 嗅探失败按标准逗号分隔（多数 csv 如此）
        rows = [[_cell(v) for v in r] for r in csv.reader(io.StringIO(text), dialect) if any(r)]
        if not rows:
            raise ParseError("CSV 里没有数据行")
        blocks = self._table_blocks("CSV", rows, doc)
        return blocks, {"pages": 0, "chars": len(text), "images": 0, "sheets": 1}

    # ==================== 标记语言 / 纯文本 ====================

    def _markdown(self, raw: bytes, doc: ParsedDocument) -> tuple[list[ParsedBlock], dict]:
        text = decode_bytes(raw)
        blocks = markdown_to_blocks(text)
        return blocks, {"pages": 0, "chars": len(text),
                       "images": sum(1 for b in blocks if b.is_media)}

    def _html(self, raw: bytes, doc: ParsedDocument) -> tuple[list[ParsedBlock], dict]:
        text = decode_bytes(raw)
        blocks = html_to_blocks(text)
        return blocks, {"pages": 0, "chars": len(text),
                       "images": sum(1 for b in blocks if b.is_media)}

    def _json(self, raw: bytes, doc: ParsedDocument) -> tuple[list[ParsedBlock], dict]:
        text = decode_bytes(raw)
        try:
            data = json.loads(text)
        except json.JSONDecodeError:
            # 解析不了的 JSON（常见于 NDJSON 日志）退化为按行文本，不因此判失败
            return markdown_to_blocks(text), {"pages": 0, "chars": len(text), "images": 0}
        blocks: list[ParsedBlock] = []
        if isinstance(data, dict):
            items: list[tuple[str, Any]] = list(data.items())
        elif isinstance(data, list):
            items = [(f"[{i}]", v) for i, v in enumerate(data)]
        else:
            items = [("", data)]
        for key, value in items:
            body = json.dumps(value, ensure_ascii=False, indent=2) \
                if isinstance(value, (dict, list)) else str(value)
            text_one = normalize_space(f"{key}: {body}" if key else str(body))
            if text_one:
                blocks.append(ParsedBlock(type=C.BLOCK_TEXT, text=text_one,
                                          meta={"json_key": key} if key else {}))
        if not blocks:
            raise ParseError("JSON 内容为空")
        return blocks, {"pages": 0, "chars": len(text), "images": 0}

    def _txt(self, raw: bytes, doc: ParsedDocument) -> tuple[list[ParsedBlock], dict]:
        text = decode_bytes(raw)
        blocks = markdown_to_blocks(text)
        return blocks, {"pages": 0, "chars": len(text), "images": 0}


def text_preview(raw: bytes, ext: str) -> str:
    """取一小段纯文本预览：auto 引擎判复杂度用，页面「原件预览」也够用，不重复走解析。"""
    if ext in C.TEXT_FORMATS:
        return normalize_space(decode_bytes(raw))[:2000]
    return ""
